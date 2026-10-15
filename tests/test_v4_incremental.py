"""v4.0 增量生成回归：仅使用临时目录中的报告副本。"""
import csv
import hashlib
import json
import os
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import openpyxl
import thermal_processor as processor


def inspect_report(path):
    workbook = openpyxl.load_workbook(path, data_only=True)
    summary = workbook['总表']
    state = {'conditions': {}, 'images': {}}
    if processor.STATE_SHEET in workbook.sheetnames:
        for kind, key, value in workbook[processor.STATE_SHEET].iter_rows(
                min_row=2, max_col=3, values_only=True):
            state[kind][key] = json.loads(value)
    result = {
        'headers': [summary.cell(4, col).value for col in range(4, 16)],
        'values': {col: [summary.cell(row, col).value for row in range(5, 28)]
                   for col in range(4, 16)},
        'channels': [summary.cell(row, 2).value for row in range(5, 28)],
        'pictures': len(summary._images),
        'anchors': [(image.anchor._from.col, image.anchor._from.row)
                    for image in summary._images],
        'charts': {s.title: len(s._charts) for s in workbook
                   if s.title not in ('总表', processor.STATE_SHEET)},
        'state': state,
    }
    workbook.close()
    return result


class IncrementalExcelTests(unittest.TestCase):
    def test_incremental_updates_and_legacy_report(self):
        import pythoncom
        import win32com.client
        with tempfile.TemporaryDirectory(prefix='thermal_v4_regression_') as directory:
            folder = Path(directory)
            report = folder / 'report.xlsx'
            shutil.copy2(ROOT / 'assets/底层模板.xlsx', report)
            config = json.loads((ROOT / 'configs/1.json').read_text(encoding='utf-8'))
            channels = [c for c in config['channels'] if c['enabled']]
            rows = processor.read_csv_rows(ROOT / '666 2026-08-06.csv')
            header, _, indices = processor.find_measurement_header(rows)
            fixture = [rows[header]] + rows[header + 1:header + 26]
            first_csv, second_csv = folder / 'first.csv', folder / 'second.csv'
            with first_csv.open('w', encoding='utf-8-sig', newline='') as target:
                csv.writer(target).writerows(fixture)
            modified = [list(row) for row in fixture]
            for row in modified[1:]:
                for index in indices.values():
                    match = processor.TEMPERATURE_RE.search(row[index])
                    if match:
                        row[index] = str(float(match.group()) + 10)
            with second_csv.open('w', encoding='utf-8-sig', newline='') as target:
                csv.writer(target).writerows(modified)
            a = {'id': 'a', 'name': '工况A', 'csv': str(first_csv), 'powerz': ''}
            b = {'id': 'b', 'name': '工况B', 'csv': str(first_csv), 'powerz': ''}
            c = {'id': 'c', 'name': '工况C', 'csv': str(first_csv), 'powerz': ''}
            i1 = {'id': 'i1', 'path': str(ROOT / 'assets/fixed_images/image3.png')}
            i2 = {'id': 'i2', 'path': str(ROOT / 'assets/fixed_images/image7.png')}
            project = {'conditions': [a, b, c], 'fixed_images': [i1, i2]}
            logs = []

            def run(conditions, images=(), channel_list=None, payload=None):
                processor.process_bulk_thermal_data(
                    str(report), conditions, images, channel_list or channels,
                    lambda _: None, logs.append, payload or project)
                return inspect_report(report)

            initial = run([a, b], [i1, i2])
            self.assertEqual(initial['headers'][:3], ['工况A', '工况B', None])
            self.assertEqual(initial['pictures'], 2)
            self.assertEqual(set(initial['anchors']), {(0, 29), (4, 29)})
            next_pass = run([c])
            self.assertEqual(next_pass['headers'][:3], ['工况A', '工况B', '工况C'])
            self.assertEqual(next_pass['values'][4], initial['values'][4])
            self.assertEqual(next_pass['values'][5], initial['values'][5])
            self.assertEqual(next_pass['pictures'], 2)

            changed_b = dict(b, csv=str(second_csv))
            updated = run([changed_b], [i2])
            self.assertEqual(updated['pictures'], 2)
            self.assertEqual(updated['values'][4], initial['values'][4])
            self.assertEqual(updated['values'][6], next_pass['values'][6])
            self.assertAlmostEqual(updated['values'][5][0] - initial['values'][5][0], 10)
            reversed_channels = run([changed_b], channel_list=list(reversed(channels)))
            self.assertEqual(reversed_channels['values'], updated['values'])
            self.assertEqual(reversed_channels['channels'], updated['channels'])

            duplicate = dict(a, id='a2')
            project['conditions'].append(duplicate)
            duplicate_pass = run([duplicate])
            self.assertEqual(duplicate_pass['headers'][:4], ['工况A', '工况B', '工况C', '工况A'])
            self.assertNotEqual(duplicate_pass['state']['conditions']['a']['sheet'],
                                duplicate_pass['state']['conditions']['a2']['sheet'])
            single_channel = run([c], channel_list=channels[1:2])
            self.assertIsNone(single_channel['values'][6][0])
            self.assertIsNotNone(single_channel['values'][6][1])
            self.assertEqual(single_channel['values'][4], initial['values'][4])

            # 手工照片占据第四格，新增照片应选择第三格，且保留手工照片。
            pythoncom.CoInitialize()
            excel = win32com.client.DispatchEx('Excel.Application')
            excel.DisplayAlerts = False
            book = None
            try:
                book = excel.Workbooks.Open(str(report), 0, False)
                sheet = book.Worksheets('总表')
                processor._add_picture_in_area(sheet, i1['path'], 'L30:O42', 'UserPhoto')
                book.Save()
            finally:
                if book is not None:
                    book.Close(False)
                excel.Quit()
                pythoncom.CoUninitialize()
            i3 = dict(i1, id='i3')
            project['fixed_images'].append(i3)
            appended = run([a], [i3])
            self.assertEqual(appended['pictures'], 4)
            self.assertEqual(appended['state']['images']['i3']['slot'], 2)
            project['fixed_images'] = [i3, i1, i2]
            reordered = run([a], [i1])
            self.assertEqual(reordered['pictures'], 4)
            self.assertEqual(reordered['state']['images']['i1']['slot'], 0)
            replacement = dict(i1, path=i2['path'])
            replaced = run([a], [replacement])
            self.assertEqual(replaced['pictures'], 4)
            self.assertEqual(replaced['state']['images']['i1']['slot'], 0)
            deletion_payload = dict(project, deleted_image_ids=['i2'])
            deleted = run([a], payload=deletion_payload)
            self.assertEqual(deleted['pictures'], 3)
            self.assertNotIn('i2', deleted['state']['images'])
            self.assertIn((11, 29), deleted['anchors'])

            # 超过容量或输入无效时，原文件必须逐字节保留。
            before = hashlib.sha256(report.read_bytes()).hexdigest()
            crowded = dict(project, conditions=project['conditions'] + [
                dict(a, id=f'extra{i}', name=f'extra{i}') for i in range(10)])
            with self.assertRaisesRegex(ValueError, '12'):
                run([a], payload=crowded)
            self.assertEqual(hashlib.sha256(report.read_bytes()).hexdigest(), before)
            with self.assertRaises(FileNotFoundError):
                run([dict(a, csv=str(folder / 'missing.csv'))])
            self.assertEqual(hashlib.sha256(report.read_bytes()).hexdigest(), before)
            extra_channels = [dict(hardware_name=hardware, custom_name=hardware)
                              for hardware in processor.parse_csv_channels(first_csv)
                              if hardware not in {item['hardware_name'] for item in channels}]
            with self.assertRaisesRegex(ValueError, '23'):
                run([a], channel_list=channels + extra_channels[:10])
            self.assertEqual(hashlib.sha256(report.read_bytes()).hexdigest(), before)

            # 构造旧版无位置记录的报告，按表头与图片内容迁移。
            legacy = folder / 'legacy.xlsx'
            shutil.copy2(ROOT / 'assets/底层模板.xlsx', legacy)
            processor.process_bulk_thermal_data(str(legacy), [a, b], [i1, i2], channels,
                                                lambda _: None, logs.append)
            pythoncom.CoInitialize()
            excel = win32com.client.DispatchEx('Excel.Application')
            excel.DisplayAlerts = False
            book = None
            try:
                book = excel.Workbooks.Open(str(legacy), 0, False)
                sheet = book.Worksheets('总表')
                for n, image in enumerate((i1, i2), 1):
                    shape_name = 'ThermalFixed_ID_' + hashlib.sha256(
                        image['id'].encode()).hexdigest()[:24]
                    sheet.Shapes(shape_name).Name = f'ThermalFixed_{n}'
                book.Worksheets(processor.STATE_SHEET).Visible = -1
                book.Worksheets(processor.STATE_SHEET).Delete()
                book.Save()
            finally:
                if book is not None:
                    book.Close(False)
                excel.Quit()
                pythoncom.CoUninitialize()
            legacy_project = {'conditions': [a, b, c], 'fixed_images': [i2, i1]}
            processor.process_bulk_thermal_data(str(legacy), [c], [i2], channels,
                                                lambda _: None, logs.append, legacy_project)
            migrated = inspect_report(legacy)
            self.assertEqual(migrated['pictures'], 2)
            self.assertEqual(migrated['state']['images']['i1']['slot'], 0)
            self.assertEqual(migrated['state']['images']['i2']['slot'], 1)
            self.assertEqual(migrated['headers'][:3], ['工况A', '工况B', '工况C'])
            self.assertTrue(all(count == 1 for count in migrated['charts'].values()))
            print('PASS: 分批、更新、换序、重名、通道变更、图片追加/替换/删除、容量、失败保留、旧报告迁移')


class GuiConfigTests(unittest.TestCase):
    def test_ids_and_pending_deletion_roundtrip(self):
        os.environ['QT_QPA_PLATFORM'] = 'offscreen'
        import thermal_gui as gui
        app = gui.QApplication.instance() or gui.QApplication([])
        with tempfile.TemporaryDirectory(prefix='thermal_gui_config_') as directory:
            base = Path(directory)
            (base / 'configs').mkdir()
            (base / 'configs/demo.json').write_text(json.dumps({
                'project_excel': str(base / 'report.xlsx'), 'channels': [],
                'conditions': [{'name': '旧工况', 'csv': '', 'powerz': ''}],
                'fixed_images': [{'path': 'old.png', 'status': '已插入'}],
            }), encoding='utf-8')
            original = gui.get_base_dir
            gui.get_base_dir = lambda: directory
            window = None
            try:
                window = gui.ThermalApp()
                cfg = window.save_project_config()
                cond_id = cfg['conditions'][0]['id']
                image_id = cfg['fixed_images'][0]['id']
                self.assertTrue(cond_id and image_id)
                window.on_project_changed()
                cfg = window.save_project_config()
                self.assertEqual(cfg['conditions'][0]['id'], cond_id)
                self.assertEqual(cfg['fixed_images'][0]['id'], image_id)
                window._queue_image_deletion(0)
                window.table_fixed_imgs.removeRow(0)
                cfg = window.save_project_config()
                self.assertEqual(cfg['deleted_image_ids'], [image_id])
                self.assertEqual(cfg['deleted_image_records'][0]['path'], 'old.png')
                window.on_project_changed()
                self.assertEqual(window.pending_image_deletions, [image_id])
            finally:
                if window:
                    window.close()
                gui.get_base_dir = original
            app.processEvents()


if __name__ == '__main__':
    unittest.main(verbosity=2)
