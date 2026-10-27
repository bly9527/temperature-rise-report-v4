# -*- coding: utf-8 -*-
"""读取温升 CSV，并让 Microsoft Excel 按参考附表生成原生图表。"""

import csv
import hashlib
import json
import math
import os
import re
import shutil
import sys
import tempfile
import zipfile
import posixpath
import xml.etree.ElementTree as ET
from datetime import datetime

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import to_excel


CHANNEL_HEADER_RE = re.compile(r'^\s*(\d{3})\s*\(')
TEMPERATURE_RE = re.compile(r'[-+]?\d*\.?\d+(?:[eE][-+]?\d+)?')
REFERENCE_NAME = '附表格式参考.xlsx'
REFERENCE_CHART_WIDTH = 569
REFERENCE_CHART_HEIGHT = 280
EXCEL_LINE_CHART = 4
EXCEL_PASTE_FORMATS = -4122
EXCEL_PICTURE = 13
EXCEL_MOVE_AND_SIZE = 1
STATE_SHEET = '_ThermalState'


def _record_id(record, kind):
    if record.get('id'):
        return str(record['id'])
    value = record.get('name', '') if kind == 'condition' else os.path.normcase(
        os.path.abspath(record['path']))
    return kind + '_' + hashlib.sha256(value.encode('utf-8')).hexdigest()[:24]


def _load_report_state(workbook):
    state = {'conditions': {}, 'images': {}}
    for sheet in workbook.Worksheets:
        if sheet.Name != STATE_SHEET:
            continue
        if sheet.Cells(1, 1).Value != 'thermal-state-v1':
            raise ValueError('报告中的工况记录格式无法识别，请使用报告备份')
        for row in range(2, sheet.UsedRange.Rows.Count + 1):
            kind, key, value = (sheet.Cells(row, col).Value for col in (1, 2, 3))
            if kind in state and key and value:
                state[kind][key] = json.loads(value)
        break
    return state


def _save_report_state(workbook, state):
    sheet = next((s for s in workbook.Worksheets if s.Name == STATE_SHEET), None)
    if sheet is None:
        sheet = workbook.Worksheets.Add()
        sheet.Name = STATE_SHEET
    sheet.Cells.ClearContents()
    sheet.Cells(1, 1).Value = 'thermal-state-v1'
    row = 2
    for kind, records in state.items():
        for key, record in records.items():
            sheet.Range(sheet.Cells(row, 1), sheet.Cells(row, 3)).Value = (
                (kind, key, json.dumps(record, ensure_ascii=False)),)
            row += 1
    sheet.Visible = 2  # xlSheetVeryHidden，随报告保存并支持复制后继续更新。


def _legacy_picture_hashes(path):
    """用图片内容匹配旧报告，避免凭本次勾选顺序替换旧图片。"""
    hashes = {}
    ns = {'xdr': 'http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing',
          'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
    relns = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if not re.fullmatch(r'xl/drawings/drawing\d+\.xml', name):
                continue
            relpath = posixpath.join(posixpath.dirname(name), '_rels',
                                    posixpath.basename(name) + '.rels')
            if relpath not in archive.namelist():
                continue
            relationships = {r.attrib['Id']: r.attrib['Target']
                             for r in ET.fromstring(archive.read(relpath))}
            for pic in ET.fromstring(archive.read(name)).findall('.//xdr:pic', ns):
                props = pic.find('.//xdr:cNvPr', ns)
                blip = pic.find('.//a:blip', ns)
                if props is None or blip is None:
                    continue
                target = relationships.get(blip.get(relns + 'embed'))
                if not target:
                    continue
                media = (target.lstrip('/') if target.startswith('/') else
                         posixpath.normpath(posixpath.join(posixpath.dirname(name), target)))
                hashes[props.get('name')] = hashlib.sha256(archive.read(media)).hexdigest()
    return hashes


def sanitize_sheet_title(title):
    safe = re.sub(r'[\/\\?*:[\]]', '_', str(title)).strip()
    return safe[:31] if safe else 'Sheet'


def read_csv_rows(csv_path):
    for enc in ('utf-8-sig', 'gb18030'):
        try:
            with open(csv_path, 'r', encoding=enc, newline='') as source:
                return list(csv.reader(source))
        except UnicodeError:
            continue
    raise ValueError(f'无法识别 CSV 文件编码：{csv_path}')


def find_measurement_header(rows):
    for row_idx, row in enumerate(rows):
        channel_indices = {}
        for col_idx, value in enumerate(row):
            match = CHANNEL_HEADER_RE.match(value.strip())
            if match:
                channel_indices[match.group(1)] = col_idx
        if not channel_indices:
            continue
        for col_idx, value in enumerate(row):
            label = value.strip().lower()
            if '扫描时间' in label or '时间' in label or 'time' in label:
                return row_idx, col_idx, channel_indices
    raise ValueError('未找到包含扫描时间和温度通道编号的测量数据表头')


def parse_csv_channels(csv_path):
    if not os.path.isfile(csv_path):
        return []
    rows = read_csv_rows(csv_path)
    header_idx, _, _ = find_measurement_header(rows)
    return [value.strip() for value in rows[header_idx]
            if CHANNEL_HEADER_RE.match(value.strip())]


def _excel_time(value):
    value = value.strip()
    try:
        # pywin32 会按本地时区转换 datetime；Excel 序列值保留 CSV 的原始时刻。
        return to_excel(datetime.fromisoformat(value))
    except ValueError:
        return value


def _scan_number(value):
    value = value.strip()
    try:
        number = float(value)
        if math.isfinite(number):
            return int(number) if number.is_integer() else number
    except ValueError:
        pass
    return value


def _prepare_condition(cond, channels_config):
    name = str(cond.get('name') or '未命名工况').strip()
    csv_path = cond.get('csv') or ''
    powerz_path = cond.get('powerz') or ''
    if not os.path.isfile(csv_path):
        raise FileNotFoundError(f'工况 [{name}] 的 CSV 文件不存在：{csv_path}')
    if powerz_path and not os.path.isfile(powerz_path):
        raise FileNotFoundError(f'工况 [{name}] 的 Power-Z 图片不存在：{powerz_path}')

    csv_rows = read_csv_rows(csv_path)
    header_idx, time_idx, channel_indices = find_measurement_header(csv_rows)
    header = csv_rows[header_idx]
    scan_idx = 1 if time_idx == 0 else 0
    selected_indices = []
    missing = []
    for channel in channels_config:
        match = re.search(r'\d+', str(channel.get('hardware_name', '')))
        number = match.group() if match else ''
        idx = channel_indices.get(number)
        if idx is None:
            missing.append(channel.get('hardware_name', ''))
        else:
            selected_indices.append(idx)
    if missing:
        raise ValueError(f'工况 [{name}] 缺少通道：{"、".join(missing)}')

    data = []
    maxima = [None] * len(channels_config)
    for row in csv_rows[header_idx + 1:]:
        if time_idx >= len(row) or not row[time_idx].strip():
            continue
        readings = []
        for idx in selected_indices:
            match = TEMPERATURE_RE.search(row[idx]) if idx < len(row) else None
            value = float(match.group()) if match else None
            readings.append(value if value is not None and math.isfinite(value) else None)
        if not any(value is not None for value in readings):
            continue
        for index, value in enumerate(readings):
            if value is not None:
                maxima[index] = value if maxima[index] is None else max(maxima[index], value)
        scan = _scan_number(row[scan_idx]) if scan_idx < len(row) else None
        data.append((_excel_time(row[time_idx]), scan, *readings))
    if not data or any(value is None for value in maxima):
        missing_data = [channels_config[i]['hardware_name'] for i, value in enumerate(maxima)
                        if value is None]
        raise ValueError(f'工况 [{name}] 没有有效温度数据：{"、".join(missing_data)}')

    return {
        'id': _record_id(cond, 'condition'),
        'name': name,
        'powerz': powerz_path,
        'header_row': header_idx + 1,
        'headers': (header[time_idx].strip(),
                    header[scan_idx].strip() if scan_idx < len(header) else '扫描编号',
                    *[header[idx].strip() for idx in selected_indices]),
        'data': data,
        'maxima': maxima,
    }


def _clean_external_links(path, log_callback):
    # 用户的纯净模板只剩外部链接缓存，且没有任何外部引用公式。
    workbook = openpyxl.load_workbook(path)
    external_formula = any(
        cell.data_type == 'f' and '[' in str(cell.value)
        for sheet in workbook for cell in sheet._cells.values()
    )
    external_name = any('[' in str(name.attr_text)
                        for name in workbook.defined_names.values())
    safe_to_remove = not external_formula and not external_name
    if workbook._external_links and safe_to_remove:
        workbook._external_links.clear()
        workbook.save(path)
        log_callback('已移除报告中无用的旧外部链接')
    workbook.close()
    return safe_to_remove


def _strip_stale_external_link_parts(path):
    """只移除无引用的外链包部件，保留 Excel 原生图表的 XML 原样。"""
    temp_fd, cleaned_path = tempfile.mkstemp(prefix='~thermal_links_', suffix='.xlsx',
                                              dir=os.path.dirname(path))
    os.close(temp_fd)
    try:
        with zipfile.ZipFile(path) as source, zipfile.ZipFile(cleaned_path, 'w') as output:
            for info in source.infolist():
                name = info.filename
                if name.startswith('xl/externalLinks/'):
                    continue
                data = source.read(name)
                if name == 'xl/workbook.xml':
                    data = re.sub(rb'<externalReferences\b[^>]*>.*?</externalReferences>',
                                  b'', data, flags=re.DOTALL)
                elif name == 'xl/_rels/workbook.xml.rels':
                    data = re.sub(
                        rb'<Relationship\b(?=[^>]*Type="[^"]*/externalLink")[^>]*/>',
                        b'', data,
                    )
                elif name == '[Content_Types].xml':
                    data = re.sub(
                        rb'<Override\b(?=[^>]*PartName="/xl/externalLinks/)[^>]*/>',
                        b'', data,
                    )
                output.writestr(info, data)
        os.replace(cleaned_path, path)
    finally:
        if os.path.exists(cleaned_path):
            os.remove(cleaned_path)


def _add_picture(sheet, image_path, cell_address, width, height, name):
    cell = sheet.Range(cell_address)
    picture = sheet.Shapes.AddPicture(os.path.abspath(image_path), 0, -1,
                                      cell.Left + 2, cell.Top + 2, width, height)
    picture.Name = name
    picture.Placement = EXCEL_MOVE_AND_SIZE


def _add_picture_in_area(sheet, image_path, area_address, name):
    """把总表截图放进对应图片格，并保持参考报告的统一长方形尺寸。"""
    area = sheet.Range(area_address)
    margin = 2
    picture = sheet.Shapes.AddPicture(os.path.abspath(image_path), 0, -1,
                                      area.Left + margin, area.Top + margin,
                                      min(285, max(area.Width - margin * 2, 1)),
                                      min(195, max(area.Height - margin * 2, 1)))
    picture.Name = name
    picture.Placement = EXCEL_MOVE_AND_SIZE


def _release_empty_condition_reservations(workbook, summary, state):
    # Older versions reserved columns for unchecked, never-generated conditions.
    # Release only empty reservations: no summary header/data and no detail sheet.
    occupied = {col for col in range(4, 16)
                if any(summary.Cells(row, col).Value is not None
                       for row in range(4, 28))}
    used_sheets = {sheet.Name.casefold() for sheet in workbook.Worksheets}
    for key, record in list(state['conditions'].items()):
        if (int(record['column']) not in occupied
                and record['sheet'].casefold() not in used_sheets):
            del state['conditions'][key]


def _assign_conditions(workbook, summary, state, registry, selected_ids):
    mappings = state['conditions']
    headers = {col: str(summary.Cells(4, col).Value or '').strip()
               for col in range(4, 16)}
    occupied = {col for col in headers if headers[col] or any(
        summary.Cells(row, col).Value is not None for row in range(5, 28))}
    used_sheets = {sheet.Name.casefold() for sheet in workbook.Worksheets}
    occupied.update(int(record['column']) for record in mappings.values())
    used_sheets.update(record['sheet'].casefold() for record in mappings.values())
    registry_ids = [_record_id(item, 'condition') for item in registry]
    if len(set(registry_ids)) != len(registry_ids):
        raise ValueError('工况编号重复；同名工况需要分别添加，不能复用同一编号')
    for item in registry:
        key = _record_id(item, 'condition')
        if key in mappings:
            continue
        name = str(item.get('name') or '未命名工况').strip()
        bound = {int(record['column']) for record in mappings.values()}
        matches = [col for col, header in headers.items() if header == name and col not in bound]
        if len(matches) > 1:
            raise ValueError(f'旧总表有多列 [{name}]，请先确认重复列，当前报告未修改')
        if matches:
            same_name = [r for r in registry if str(r.get('name', '')).strip() == name]
            inserted = [r for r in same_name if r.get('status') == '已插入']
            if len(same_name) > 1 and (len(inserted) != 1 or inserted[0] is not item):
                # 已有工况先绑定，新增同名工况随后分配空列。
                if len(inserted) == 1:
                    continue
                raise ValueError(f'旧报告工况 [{name}] 重名，无法确定所属工况，请先核对名称')
            col = matches[0]
            sheet_name = sanitize_sheet_title(name)
            existing = [s.Name for s in workbook.Worksheets
                        if s.Name.casefold() == sheet_name.casefold()]
            if existing:
                if any(record['sheet'].casefold() == existing[0].casefold()
                       for record in mappings.values()):
                    raise ValueError(f'旧报告工况 [{name}] 的附表名与其他工况冲突，请先核对名称')
                sheet_name = existing[0]
            else:
                sheet_name = None
        else:
            if key not in selected_ids:
                continue
            col = next((c for c in range(4, 16) if c not in occupied), None)
            if col is None:
                raise ValueError('总表最多容纳 12 个工况；已有工况也计入，不能覆盖“规格/判定”列')
            sheet_name = None
        if sheet_name is None:
            base = sanitize_sheet_title(name)
            sheet_name = base
            number = 1
            while sheet_name.casefold() in used_sheets:
                suffix = f' ({number})'
                sheet_name = base[:31 - len(suffix)] + suffix
                number += 1
        mappings[key] = {'name': name, 'column': col, 'sheet': sheet_name}
        occupied.add(col)
        used_sheets.add(sheet_name.casefold())
    # 新增同名工况在已有工况完成旧列绑定后再分配。
    pending = [item for item in registry
               if _record_id(item, 'condition') in selected_ids
               and _record_id(item, 'condition') not in mappings]
    if pending:
        _assign_conditions(workbook, summary, state, pending, selected_ids)


def _validate_condition_positions(state):
    records = list(state['conditions'].values())
    columns = [int(record['column']) for record in records]
    sheets = [record['sheet'].casefold() for record in records]
    if (len(columns) != len(set(columns)) or any(c < 4 or c > 15 for c in columns)
            or len(sheets) != len(set(sheets))
            or any(s in ('总表', STATE_SHEET.casefold()) for s in sheets)):
        raise ValueError('报告工况位置记录异常，请先核对报告；原报告未修改')


def _compact_condition_columns(summary, state):
    """Move only the condition cells, leaving specification columns and photos alone."""
    occupied = {int(r['column']) for r in state['conditions'].values()}
    occupied.update(col for col in range(4, 16)
                    if any(summary.Cells(row, col).Value is not None
                           for row in range(4, 28)))
    moves = {source: target for target, source in enumerate(sorted(occupied), 4)}
    for source, target in moves.items():
        if source == target:
            continue
        old = summary.Range(summary.Cells(4, source), summary.Cells(27, source))
        new = summary.Range(summary.Cells(4, target), summary.Cells(27, target))
        old.Copy(new)
        old.ClearContents()
    for record in state['conditions'].values():
        record['column'] = moves[int(record['column'])]


def _order_condition_sheets(workbook, summary, state):
    sheets = {s.Name.casefold(): s for s in workbook.Worksheets}
    previous = summary
    for record in sorted(state['conditions'].values(), key=lambda r: int(r['column'])):
        sheet = sheets.get(record['sheet'].casefold())
        if sheet is not None:
            if sheet.Index != previous.Index + 1:
                sheet.Move(None, previous)
            previous = sheet


def _condition_result_to_delete(summary, state, condition, registry):
    key = _record_id(condition, 'condition')
    if key in state['conditions']:
        record = state['conditions'][key]
        header = str(summary.Cells(4, int(record['column'])).Value or '').strip()
        if header and header != record['name']:
            raise ValueError('该工况的总表位置与记录不一致，请先核对报告；原报告未修改')
        return key, record
    name = str(condition.get('name') or '').strip()
    if not name and condition.get('status') != '已插入':
        return key, None
    matches = [col for col in range(4, 16)
               if str(summary.Cells(4, col).Value or '').strip() == name]
    inserted = condition.get('status') == '已插入'
    if not matches:
        if inserted:
            raise ValueError('未找到该工况对应的报告结果，请先核对报告；原报告未修改')
        return key, None
    owners = [r for r in state['conditions'].values() if int(r['column']) in matches]
    if owners:
        if inserted:
            raise ValueError('同名结果属于其他工况，无法安全删除；原报告未修改')
        return key, None  # An unchecked new namesake must not delete an existing result.
    same_name = [r for r in registry if str(r.get('name') or '').strip() == name]
    completed = [r for r in same_name if r.get('status') == '已插入']
    if not inserted and completed:
        return key, None
    if len(matches) != 1 or (len(same_name) > 1 and not (
            len(completed) == 1 and _record_id(completed[0], 'condition') == key)):
        raise ValueError('旧报告存在重名工况，无法确定要删除的结果；原报告未修改')
    sheet_name = sanitize_sheet_title(name)
    if any(r['sheet'].casefold() == sheet_name.casefold()
           for r in state['conditions'].values()):
        raise ValueError('对应附表属于其他工况，无法安全删除；原报告未修改')
    return key, {'name': name, 'column': matches[0], 'sheet': sheet_name}


def _commit_condition_deletion(work_path, report_path, config_path, new_config):
    """Stage both files; roll the workbook back if committing the config fails."""
    fd, config_temp = tempfile.mkstemp(prefix='~condition_', suffix='.json',
                                      dir=os.path.dirname(os.path.abspath(config_path)))
    backup = None
    try:
        with os.fdopen(fd, 'w', encoding='utf-8') as stream:
            json.dump(new_config, stream, ensure_ascii=False, indent=2)
        if work_path is not None:
            fd, backup = tempfile.mkstemp(prefix='~condition_backup_', suffix='.xlsx',
                                          dir=os.path.dirname(report_path))
            os.close(fd)
            shutil.copy2(report_path, backup)
            try:
                os.replace(work_path, report_path)
            except PermissionError as error:
                raise PermissionError('请先关闭 Excel/WPS 中的工程报告，再重试删除。') from error
        try:
            os.replace(config_temp, config_path)
        except OSError as error:
            if backup is not None:
                try:
                    os.replace(backup, report_path)
                except OSError as rollback_error:
                    recovery = backup
                    backup = None  # Keep the recovery copy if another app locked the report.
                    raise RuntimeError(f'配置保存失败，报告备份保存在：{recovery}') from rollback_error
            raise OSError('工程配置保存失败，删除未完成；请检查配置文件是否被占用。') from error
    finally:
        for path in (config_temp, backup):
            if path and os.path.exists(path):
                os.remove(path)


def delete_condition_result(project_excel, condition, project_config, config_path,
                            log_callback=None):
    """Immediately remove one condition's output and persist its remaining config."""
    log = log_callback or (lambda message: None)
    report_path = os.path.abspath(project_excel)
    if not os.path.isfile(report_path):
        raise FileNotFoundError('工程报告不存在，请先核对报告路径；未删除工况。')
    with open(report_path, 'rb') as stream:
        original_hash = hashlib.sha256(stream.read()).digest()
    registry = project_config.get('conditions', [])
    key = _record_id(condition, 'condition')
    if sum(_record_id(c, 'condition') == key for c in registry) != 1:
        raise ValueError('工况编号异常，无法确定删除对象；原报告未修改')
    new_config = dict(project_config)
    new_config['conditions'] = [c for c in registry if _record_id(c, 'condition') != key]
    fd, work_path = tempfile.mkstemp(prefix='~thermal_delete_', suffix='.xlsx',
                                    dir=os.path.dirname(report_path))
    os.close(fd)
    try:
        shutil.copy2(report_path, work_path)
        import pythoncom
        pythoncom.CoInitialize()
        excel = book = None
        changed = False
        try:
            excel = _create_spreadsheet_application(log)
            excel.Visible = False
            excel.DisplayAlerts = False
            excel.AskToUpdateLinks = False
            excel.EnableEvents = False
            book = excel.Workbooks.Open(work_path, 0, False)
            if book.ReadOnly:
                raise PermissionError('报告无法编辑，请关闭 Excel/WPS 中的报告后重试。')
            summary = book.Worksheets.Item('总表')
            state = _load_report_state(book)
            _validate_condition_positions(state)
            key, record = _condition_result_to_delete(summary, state, condition, registry)
            if record is not None:
                sheet_name = record['sheet']
                if sheet_name.casefold() in ('总表', STATE_SHEET.casefold()):
                    raise ValueError('工况附表记录异常，不能删除总表；原报告未修改')
                for sheet in book.Worksheets:
                    if sheet.Name.casefold() == sheet_name.casefold():
                        sheet.Delete()
                        break
                col = int(record['column'])
                summary.Range(summary.Cells(4, col), summary.Cells(27, col)).ClearContents()
                state['conditions'].pop(key, None)
                _release_empty_condition_reservations(book, summary, state)
                _compact_condition_columns(summary, state)
                excel.CutCopyMode = False
                _order_condition_sheets(book, summary, state)
                _save_report_state(book, state)
                summary.Activate()
                book.Save()
                changed = True
            book.Close(False)
            book = None
        finally:
            try:
                if book is not None:
                    book.Close(False)
            finally:
                try:
                    if excel is not None:
                        excel.Quit()
                finally:
                    pythoncom.CoUninitialize()
        with open(report_path, 'rb') as stream:
            if hashlib.sha256(stream.read()).digest() != original_hash:
                raise ValueError('处理期间报告已被其他程序修改，请重新检查后再删除。')
        _commit_condition_deletion(work_path if changed else None, report_path,
                                   config_path, new_config)
        log('已删除工况结果并整理总表排列。' if changed else '该工况尚无生成结果，已从工程配置移除。')
        return new_config
    finally:
        if os.path.exists(work_path):
            os.remove(work_path)


def _write_summary(summary, prepared, channels_config, state):
    channel_rows = {}
    for row in range(5, 28):
        value = summary.Cells(row, 2).Value
        match = re.search(r'\d+', str(value or ''))
        if match:
            key = match.group()
            if key in channel_rows:
                raise ValueError(f'总表通道 [{key}] 重复，无法确定最高温应填写的行')
            channel_rows[key] = row
    rows = []
    for channel in channels_config:
        key = re.search(r'\d+', channel['hardware_name']).group()
        row = channel_rows.get(key)
        if row is None:
            row = next((r for r in range(5, 28) if not summary.Cells(r, 2).Value
                        and all(summary.Cells(r, c).Value is None for c in range(4, 16))), None)
            if row is None:
                raise ValueError('总表测温通道格已满（最多23行），不能覆盖图片区域')
            channel_rows[key] = row
        rows.append(row)
        summary.Cells(row, 1).Value = row - 4
        summary.Cells(row, 2).Value = channel['hardware_name']
        summary.Cells(row, 3).Value = channel['custom_name'].strip()
    for cond in prepared:
        record = state['conditions'][cond['id']]
        col = record['column']
        cond['sheet_name'] = record['sheet']
        record['name'] = cond['name']
        summary.Cells(4, col).Value = cond['name']
        # 只重写当前工况，未选择通道留空，避免混入同列的旧数据。
        summary.Range(summary.Cells(5, col), summary.Cells(27, col)).ClearContents()
        for row, maximum in zip(rows, cond['maxima']):
            cell = summary.Cells(row, col)
            cell.Value = round(maximum, 2)
            cell.NumberFormat = '0.00'


def _picture_area(slot):
    columns = (('A', 'D'), ('E', 'H'), ('I', 'K'), ('L', 'O'))
    left, right = columns[slot % 4]
    picture_row = slot // 4
    first = 30 if picture_row == 0 else 43 + (picture_row - 1) * 14
    last = 42 if picture_row == 0 else first + 13
    return f'{left}{first}:{right}{last}'


def _excel_legacy_image_hashes(workbook, summary, state, registry, work_path):
    """用同一 Excel 的保存结果核对旧图，避免压缩或编码差异造成误判。"""
    bound = {record['shape'] for record in state['images'].values()}
    old_shapes = [shape for shape in summary.Shapes
                  if shape.Type == EXCEL_PICTURE and shape.Name not in bound
                  and re.fullmatch(r'ThermalFixed_\d+', shape.Name)]
    candidates = [image for image in registry
                  if _record_id(image, 'image') not in state['images']
                  and os.path.isfile(image['path'])]
    if not old_shapes or not candidates:
        return _legacy_picture_hashes(work_path), {}
    sizes = {(round(shape.Width, 3), round(shape.Height, 3)) for shape in old_shapes}
    probe = workbook.Worksheets.Add()
    expected_names = {}
    try:
        index = 0
        for image in candidates:
            key = _record_id(image, 'image')
            for width, height in sizes:
                name = f'ThermalImageProbe_{index}'
                picture = probe.Shapes.AddPicture(os.path.abspath(image['path']), 0, -1,
                                                   0, 0, width, height)
                picture.Name = name
                expected_names[name] = key
                index += 1
        workbook.Save()
        hashes = _legacy_picture_hashes(work_path)
        expected = {}
        for name, key in expected_names.items():
            if name in hashes:
                expected.setdefault(key, set()).add(hashes[name])
        return hashes, expected
    finally:
        summary.Activate()
        probe.Delete()


def _write_summary_images(summary, state, selected, registry, deleted, legacy_hashes, log,
                          expected_hashes=None):
    records = state['images']
    shapes = {s.Name: s for s in summary.Shapes if s.Type == EXCEL_PICTURE}
    registered_hashes = {}
    for image in registry:
        key = _record_id(image, 'image')
        if key not in records and os.path.isfile(image['path']):
            with open(image['path'], 'rb') as source:
                digest = hashlib.sha256(source.read()).hexdigest()
            digests = {digest} | (expected_hashes or {}).get(key, set())
            for value in digests:
                candidates = registered_hashes.setdefault(value, [])
                if not any(candidate[0] == key for candidate in candidates):
                    candidates.append((key, image))
    bound_names = {r['shape'] for r in records.values()}
    for name, shape in shapes.items():
        match = re.fullmatch(r'ThermalFixed_(\d+)', name)
        if name in bound_names or not match:
            continue
        candidates = registered_hashes.get(legacy_hashes.get(name), [])
        if len(candidates) == 1:
            key, image = candidates[0]
            if key not in records:
                records[key] = {'shape': name, 'slot': int(match.group(1)) - 1,
                                'path': image['path']}
                bound_names.add(name)
                log('已根据图片内容识别旧报告图片，保留原图片格')
    if any(re.fullmatch(r'ThermalFixed_\d+', name) and name not in bound_names
           for name in shapes):
        log('部分旧图片无法确定对应关系，已保留；新图将使用空图片格')
    for key in deleted:
        legacy_record = next((image for image in registry
                              if _record_id(image, 'image') == key), None)
        if key not in records and legacy_record and legacy_record.get('status') == '已插入':
            raise ValueError('无法识别待删除的旧图片，请先在 Excel 中核对；当前报告未修改')
        record = records.pop(key, None)
        if record and record['shape'] in shapes:
            shapes.pop(record['shape']).Delete()
    used_slots = {int(r['slot']) for r in records.values()}
    managed_names = {record['shape'] for record in records.values()}
    for name, shape in shapes.items():
        if name in managed_names:
            continue
        match = re.fullmatch(r'ThermalFixed_(\d+)', name)
        if match:
            used_slots.add(int(match.group(1)) - 1)
        # 用户手动插入的照片也占格，但绝不自动删除。
        elif shape.BottomRightCell.Row >= 30 and shape.TopLeftCell.Column <= 15:
            last_row = shape.BottomRightCell.Row
            last_group = 0 if last_row <= 42 else 1 + (last_row - 43) // 14
            for slot in range((last_group + 1) * 4):
                area = summary.Range(_picture_area(slot))
                overlap_width = min(shape.Left + shape.Width, area.Left + area.Width) - max(shape.Left, area.Left)
                overlap_height = min(shape.Top + shape.Height, area.Top + area.Height) - max(shape.Top, area.Top)
                # Excel 保存时锚点有小幅舍入，不把边框处的偏差算作占据相邻格。
                if overlap_width > 4 and overlap_height > 4:
                    used_slots.add(slot)
    for image in selected:
        key = _record_id(image, 'image')
        record = records.get(key)
        if record:
            slot, name = int(record['slot']), record['shape']
            if name in shapes:
                shapes.pop(name).Delete()
        else:
            slot = next(i for i in range(len(used_slots) + 1) if i not in used_slots)
            name = 'ThermalFixed_ID_' + hashlib.sha256(key.encode('utf-8')).hexdigest()[:24]
        _add_picture_in_area(summary, image['path'], _picture_area(slot), name)
        records[key] = {'shape': name, 'slot': slot, 'path': image['path']}
        used_slots.add(slot)


def _excel_environment_description(excel):
    parts = []
    for label, attribute in (('程序', 'Name'), ('版本', 'Version'), ('构建', 'Build')):
        try:
            parts.append(f'{label}={getattr(excel, attribute)}')
        except Exception:
            pass
    return '，'.join(parts) or '表格程序版本未知'


def _create_spreadsheet_application(log_callback):
    """优先使用 Excel；没有可启动的 Excel 时使用 WPS 表格。"""
    import pywintypes
    import win32com.client

    failures = []
    for progid, label in (('Excel.Application', 'Microsoft Excel'),
                          ('ket.Application', 'WPS 表格')):
        try:
            application = win32com.client.DispatchEx(progid)
        except pywintypes.com_error as error:
            failures.append(f'{label}：{error}')
            continue
        log_callback(f'已启动表格程序：{label}（{progid}）')
        return application
    raise RuntimeError('无法启动 Excel 或 WPS 表格，请确认已安装其中一种并能正常打开。'
                       + '；'.join(failures))


def _insert_native_line_chart(sheet, cond, count, last_letter, last_row,
                              environment, log_callback):
    """新接口失败时改用传统原生图表接口，数据引用保持一致。"""
    import pywintypes

    before_count = sheet.ChartObjects().Count
    try:
        sheet.Range(f'A1:A{last_row},C1:{last_letter}{last_row}').Select()
        container = sheet.Shapes.AddChart2(-1, EXCEL_LINE_CHART)
        chart = container.Chart
        if chart.SeriesCollection().Count != count:
            raise ValueError('自动识别的温度系列数量不正确')
        for series_idx in range(1, count + 1):
            chart.SeriesCollection(series_idx).XValues = sheet.Range(f'A3:A{last_row}')
    except (AttributeError, ValueError, pywintypes.com_error) as modern_error:
        # 部分程序在报错前已创建空图表，只清理本次调用产生的对象。
        while sheet.ChartObjects().Count > before_count:
            sheet.ChartObjects().Item(sheet.ChartObjects().Count).Delete()
        log_callback(f'[{cond["name"]}] AddChart2 不可用，切换原生兼容接口；'
                     f'{environment}；原因：{modern_error}')
        try:
            container = sheet.ChartObjects().Add(
                sheet.Range('E9').Left, sheet.Range('E9').Top,
                REFERENCE_CHART_WIDTH, REFERENCE_CHART_HEIGHT)
            chart = container.Chart
            chart.ChartType = EXCEL_LINE_CHART
            while chart.SeriesCollection().Count:
                chart.SeriesCollection(1).Delete()
            sheet_ref = "'" + sheet.Name.replace("'", "''") + "'!"
            for series_idx in range(1, count + 1):
                column = get_column_letter(series_idx + 2)
                series = chart.SeriesCollection().NewSeries()
                series.Formula = (
                    f'=SERIES({sheet_ref}${column}$1:${column}$2,'
                    f'{sheet_ref}$A$3:$A${last_row},'
                    f'{sheet_ref}${column}$3:${column}${last_row},{series_idx})')
            chart.HasLegend = count > 1
            if chart.HasLegend:
                chart.Legend.Position = -4107  # xlLegendPositionBottom
            log_callback(f'[{cond["name"]}] 已通过 ChartObjects.Add 插入原生折线图')
        except Exception as legacy_error:
            raise RuntimeError(
                f'工况 [{cond["name"]}] 无法创建原生折线图。{environment}。'
                f'AddChart2：{modern_error}；ChartObjects.Add：{legacy_error}。'
                '请提供本机 Excel/WPS 的名称与版本，以便继续排查。') from legacy_error

    container.Left = sheet.Range('E9').Left
    container.Top = sheet.Range('E9').Top
    container.Width = REFERENCE_CHART_WIDTH
    container.Height = REFERENCE_CHART_HEIGHT
    chart.HasTitle = True
    chart.ChartTitle.Text = cond['name']
    container.Placement = EXCEL_MOVE_AND_SIZE


def _write_detail(excel, workbook, reference_sheet, cond, channels_config,
                  log_callback=None):
    name = cond['sheet_name']
    for sheet in workbook.Worksheets:
        if sheet.Name == name:
            sheet.Delete()
            break
    # pywin32 的动态 COM 调用需按位置传递 After；仅用关键字会另建工作簿。
    reference_sheet.Copy(None, workbook.Worksheets(workbook.Worksheets.Count))
    sheet = workbook.Worksheets(workbook.Worksheets.Count)
    sheet.Name = name
    for chart_idx in range(sheet.ChartObjects().Count, 0, -1):
        sheet.ChartObjects().Item(chart_idx).Delete()

    count = len(channels_config)
    last_col = count + 2  # A=扫描时间，B=扫描编号，C起为温度通道。
    last_letter = get_column_letter(last_col)
    last_row = len(cond['data']) + 3
    sheet.Range('A1:P829').ClearContents()

    # Excel 跨工作簿复制工作表不会完整保留列宽和默认行高。
    for col in range(1, min(last_col, 16) + 1):
        sheet.Columns(col).ColumnWidth = reference_sheet.Columns(col).ColumnWidth
    sheet.Rows(f'1:{last_row}').RowHeight = reference_sheet.Rows(4).RowHeight

    # 参考文件的第 4 行至 829 行样式一致，超出部分续用第 4 行格式。
    for col in range(17, last_col + 1):
        letter = get_column_letter(col)
        sheet.Range('P1:P829').Copy()
        sheet.Range(f'{letter}1:{letter}829').PasteSpecial(EXCEL_PASTE_FORMATS)
        sheet.Columns(col).ColumnWidth = reference_sheet.Columns(16).ColumnWidth
    if last_row > 829:
        sheet.Range(f'A4:{last_letter}4').Copy()
        sheet.Range(f'A830:{last_letter}{last_row}').PasteSpecial(EXCEL_PASTE_FORMATS)
    excel.CutCopyMode = False

    sheet.Range(sheet.Cells(1, 3), sheet.Cells(1, last_col)).Value = (
        tuple(channel['custom_name'].strip() for channel in channels_config),
    )
    sheet.Range(sheet.Cells(3, 1), sheet.Cells(3, last_col)).Value = (
        cond['headers'],
    )
    sheet.Range(sheet.Cells(4, 1), sheet.Cells(last_row, last_col)).Value = tuple(cond['data'])
    sheet.Range(sheet.Cells(2, 3), sheet.Cells(2, last_col)).Formula = (
        tuple(f'=MAX({get_column_letter(col)}4:{get_column_letter(col)}{last_row})'
              for col in range(3, last_col + 1)),
    )
    sheet.Range(f'A4:A{last_row}').NumberFormat = reference_sheet.Range('A4').NumberFormat

    # 全量选中扫描时间与温度通道，跳过 B 列扫描编号，交给 Excel 原生插入折线图。
    workbook.Activate()
    sheet.Activate()
    excel.ActiveWindow.Zoom = 55
    _insert_native_line_chart(sheet, cond, count, last_letter, last_row,
                              _excel_environment_description(excel),
                              log_callback or (lambda message: None))

    if cond['powerz']:
        # 14 个及以下通道保持 Q1；更多通道时始终紧接最后一个温度列。
        powerz_col = max(17, last_col + 1)
        _add_picture(sheet, cond['powerz'],
                     f'{get_column_letter(powerz_col)}1',
                     285, 210, 'ThermalPowerZ')
    return last_row


def process_bulk_thermal_data(project_excel, condition_configs, fixed_images,
                              channels_config, progress_callback, log_callback,
                              project_config=None):
    if not os.path.isfile(project_excel):
        raise FileNotFoundError(f'工程 Excel 文件不存在：{project_excel}')
    if not condition_configs or not channels_config:
        raise ValueError('至少需要一个工况和一个测温通道')
    channel_ids = [re.search(r'\d+', channel['hardware_name']).group()
                   for channel in channels_config]
    if len(channel_ids) != len(set(channel_ids)):
        raise ValueError('测温通道编号重复，请检查通道列表')
    selected_images = [dict(image) if isinstance(image, dict) else {'path': image}
                       for image in fixed_images]
    for image in selected_images:
        if not os.path.isfile(image['path']):
            raise FileNotFoundError(f'固定图片不存在：{image["path"]}')
    image_ids = [_record_id(image, 'image') for image in selected_images]
    if len(image_ids) != len(set(image_ids)):
        raise ValueError('本次图片编号重复，请检查图片列表')
    project_config = project_config or {}
    registry = project_config.get('conditions', condition_configs)
    image_registry = project_config.get('fixed_images', selected_images)
    image_registry = list(image_registry) + project_config.get('deleted_image_records', [])
    image_registry = [dict(image) if isinstance(image, dict) else {'path': image}
                      for image in image_registry]

    prepared = []
    for cond in condition_configs:
        item = _prepare_condition(cond, channels_config)
        prepared.append(item)
        log_callback(f'[{item["name"]}] CSV 第 {item["header_row"] + 1} 行开始，'
                     f'{len(item["data"])} 行数据，{len(channels_config)} 个通道已匹配')

    resource_dir = getattr(sys, '_MEIPASS', os.path.dirname(__file__))
    reference_path = os.path.join(resource_dir, 'assets', REFERENCE_NAME)
    if not os.path.isfile(reference_path):
        raise FileNotFoundError(f'缺少附表格式参考文件：{reference_path}')

    report_path = os.path.abspath(project_excel)
    temp_fd, work_path = tempfile.mkstemp(prefix='~thermal_', suffix='.xlsx',
                                          dir=os.path.dirname(report_path))
    os.close(temp_fd)
    try:
        shutil.copy2(report_path, work_path)
        may_strip_links = _clean_external_links(work_path, log_callback)
        legacy_hashes = _legacy_picture_hashes(work_path)

        import pythoncom
        import win32com.client
        pythoncom.CoInitialize()
        excel = None
        output_book = None
        reference_book = None
        try:
            excel = _create_spreadsheet_application(log_callback)
            excel.Visible = False
            excel.DisplayAlerts = False
            excel.AskToUpdateLinks = False
            excel.EnableEvents = False
            log_callback(f'表格自动化环境：{_excel_environment_description(excel)}')
            output_book = excel.Workbooks.Open(work_path, 0, False)
            if output_book.ReadOnly:
                raise PermissionError(f'工程 Excel 正在被占用：{report_path}')
            reference_book = excel.Workbooks.Open(os.path.abspath(reference_path), 0, True)
            reference_sheet = reference_book.Worksheets.Item(1)
            summary = output_book.Worksheets.Item('总表')
            state = _load_report_state(output_book)
            _validate_condition_positions(state)
            _release_empty_condition_reservations(output_book, summary, state)
            _compact_condition_columns(summary, state)
            excel.CutCopyMode = False
            _assign_conditions(output_book, summary, state, registry,
                               {item['id'] for item in prepared})
            _write_summary(summary, prepared, channels_config, state)
            legacy_hashes, expected_hashes = _excel_legacy_image_hashes(
                output_book, summary, state, image_registry, work_path)
            _write_summary_images(summary, state, selected_images, image_registry,
                                  project_config.get('deleted_image_ids', []),
                                  legacy_hashes, log_callback, expected_hashes)
            for index, cond in enumerate(prepared):
                last_row = _write_detail(excel, output_book, reference_sheet,
                                         cond, channels_config, log_callback)
                log_callback(f'[{cond["name"]}] Excel 原生折线图已插入，'
                             f'扫描时间 A3:A{last_row}，温度 C1:'
                             f'{get_column_letter(len(channels_config) + 2)}{last_row}')
                progress_callback(int((index + 1) / len(prepared) * 100))
            _order_condition_sheets(output_book, summary, state)
            _save_report_state(output_book, state)
            summary.Activate()
            output_book.Save()
            output_book.Close(False)
            output_book = None
            reference_book.Close(False)
            reference_book = None
        finally:
            if output_book is not None:
                output_book.Close(False)
            if reference_book is not None:
                reference_book.Close(False)
            if excel is not None:
                excel.Quit()
            pythoncom.CoUninitialize()

        if may_strip_links:
            _strip_stale_external_link_parts(work_path)

        try:
            os.replace(work_path, report_path)
        except PermissionError as error:
            raise PermissionError(
                f'无法替换报告文件，请先关闭 Excel/WPS 中的 {os.path.basename(report_path)}'
            ) from error
        log_callback('报告已由 Excel 原生图表流程生成完毕')
    finally:
        if os.path.exists(work_path):
            os.remove(work_path)
